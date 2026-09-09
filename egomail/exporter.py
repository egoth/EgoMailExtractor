from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .extraction import compute_fields, parse_money, record_key
from .profile import normalize_field_type


def field_names(profile: dict[str, Any]) -> list[str]:
    fields = profile.get("fields", [])
    names = []
    for f in fields:
        if isinstance(f, str):
            names.append(f)
        elif isinstance(f, dict) and f.get("name"):
            names.append(f["name"])
    return names


def field_map(profile: dict[str, Any]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for f in profile.get("fields", []):
        if isinstance(f, str):
            out[f] = {"name": f, "type": "text"}
        elif isinstance(f, dict) and f.get("name"):
            item = dict(f)
            item["type"] = normalize_field_type(item.get("type"))
            out[item["name"]] = item
    return out




def _strict_number(value: Any) -> float | int | None:
    """Converte solo stringhe che sono realmente numeriche.

    Evita il vecchio comportamento pericoloso in cui un codice alfanumerico
    come ``HM895ZFM4S`` diventava ``8954`` perché venivano eliminati tutti i
    caratteri non numerici prima della conversione.
    """
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    if value in (None, ""):
        return None
    text = str(value).replace("\xa0", " ").strip()
    # Sono ammessi suffissi valuta, spazi, segno, separatore migliaia e decimali.
    text = re.sub(r"\s*(?:EUR|€)\s*$", "", text, flags=re.I).strip()
    numeric_re = re.compile(
        r"^[+-]?(?:(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?|\d+(?:\.\d+)?)$"
    )
    if not numeric_re.fullmatch(text):
        return None
    parsed = parse_money(text)
    if parsed is None:
        return None
    if re.fullmatch(r"[+-]?\d+", text):
        try:
            return int(text)
        except ValueError:
            pass
    return parsed

def excel_value(value: Any, field: dict[str, Any]) -> Any:
    """Converte il valore nel vero tipo cella richiesto dal profilo."""
    if value in (None, ""):
        return None
    kind = normalize_field_type(field.get("type"))
    if kind == "text":
        # Forza il testo, utile anche per codici con zeri iniziali.
        return str(value)
    if kind == "number":
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (int, float)):
            return value
        parsed = _strict_number(value)
        # Se il campo è stato configurato erroneamente come Numero ma contiene
        # un codice alfanumerico, non distruggiamo il dato: lo lasciamo testo.
        return parsed if parsed is not None else str(value)
    if kind == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y"):
            try:
                return datetime.strptime(text[:10], fmt).date()
            except Exception:
                pass
        try:
            return date.fromisoformat(text[:10])
        except Exception:
            return value
    if kind == "datetime":
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        text = str(value).strip().replace(",", " ")
        text = re.sub(r"\s+", " ", text)
        try:
            return datetime.fromisoformat(text)
        except Exception:
            pass
        for fmt in (
            "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S.%f",
            "%m/%d/%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M:%S.%f",
        ):
            try:
                return datetime.strptime(text, fmt)
            except Exception:
                pass
        return value
    return value


def final_formula_aliases(profile: dict[str, Any]) -> dict[str, str]:
    """Alias simbolici disponibili nelle formule riepilogative nominate.

    Esempi: ``somma_importo`` e ``conta_nome_ospite``. Il valore restituito
    è il nome del campo; il riferimento alla cella viene determinato in fase
    di scrittura perché dipende dall'ordine corrente delle colonne.
    """
    aliases: dict[str, str] = {}
    for field in profile.get("fields", []) or []:
        if not isinstance(field, dict) or not field.get("name"):
            continue
        name = str(field["name"])
        kind = str(field.get("final_formula", "") or "").lower()
        if kind == "sum":
            aliases[f"somma_{name}"] = name
        elif kind == "count":
            aliases[f"conta_{name}"] = name
    return aliases



def summary_formula_aliases(profile: dict[str, Any]) -> dict[str, str]:
    """Alias utilizzabili per riferire altre formule riepilogative.

    Il nome visuale viene trasformato in un identificatore stabile, ad esempio
    ``pagato mese corrente`` -> ``pagato_mese_corrente``. Gli alias duplicati
    vengono ignorati per evitare riferimenti ambigui.
    """
    candidates: list[tuple[str, str]] = []
    for item in profile.get("summary_formulas", []) or []:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name", "") or "").strip()
        if not name:
            continue
        alias = re.sub(r"[^\w]+", "_", name, flags=re.UNICODE).strip("_")
        if not alias:
            continue
        if alias[0].isdigit():
            alias = "formula_" + alias
        candidates.append((alias, name))
    counts: dict[str, int] = {}
    for alias, _name in candidates:
        counts[alias.casefold()] = counts.get(alias.casefold(), 0) + 1
    return {
        alias: name for alias, name in candidates
        if counts.get(alias.casefold(), 0) == 1
    }


def expand_summary_formula(
    formula: str,
    profile: dict[str, Any],
    alias_cells: dict[str, str],
) -> str:
    """Converte alias di campi finali e formule riepilogative in riferimenti Excel."""
    text = str(formula or "").strip()
    if not text:
        raise ValueError("Formula riepilogativa vuota")
    if not text.startswith("="):
        text = "=" + text

    expanded = text
    # Sostituzione per alias completi, dal più lungo al più corto.
    for alias, cell in sorted(alias_cells.items(), key=lambda kv: -len(kv[0])):
        expanded = re.sub(
            rf"(?<!\w){re.escape(alias)}(?!\w)",
            str(cell),
            expanded,
            flags=re.I,
        )

    # Manteniamo il controllo esplicito sugli alias somma_/conta_ inesistenti.
    known = {k.casefold() for k in alias_cells}
    unknown = {
        m.group(0) for m in re.finditer(
            r"\b(?:somma|conta)_[A-Za-z0-9_À-ÖØ-öø-ÿ]+\b", text, flags=re.I
        )
        if m.group(0).casefold() not in known
    }
    if unknown:
        valid = ", ".join(sorted(alias_cells)) or "<nessun alias disponibile>"
        raise ValueError(
            "Alias formula non definiti: " + ", ".join(sorted(unknown)) +
            ". Alias disponibili: " + valid
        )
    return expanded


def _is_generated_summary_row(
    ws,
    row_idx: int,
    fields: list[str],
    definitions: dict[str, dict[str, Any]],
    profile: dict[str, Any] | None = None,
) -> bool:
    """Riconosce righe riepilogative generate da EgoMailExtractor."""
    # Riga unica delle formule finali associate ai campi.
    for col_idx, field_name in enumerate(fields, 1):
        kind = str(definitions.get(field_name, {}).get("final_formula", "") or "").lower()
        if kind not in {"sum", "count"}:
            continue
        value = ws.cell(row_idx, col_idx).value
        if not isinstance(value, str) or not value.startswith("="):
            continue
        upper = value.upper()
        if kind == "sum" and upper.startswith("=SUM("):
            return True
        if kind == "count" and upper.startswith("=COUNTA("):
            return True

    # Righe delle formule riepilogative nominate: nome in colonna A e formula
    # Excel generata in colonna B.
    if profile:
        names = {
            str(item.get("name", "") or "").strip().casefold()
            for item in profile.get("summary_formulas", []) or []
            if isinstance(item, dict) and item.get("name")
        }
        if names:
            label = str(ws.cell(row_idx, 1).value or "").strip().casefold()
            formula = ws.cell(row_idx, 2).value
            if label in names and isinstance(formula, str) and formula.startswith("="):
                return True
    return False

def load_existing_records(path: str | Path, profile: dict[str, Any]) -> list[dict[str, Any]]:
    p = Path(path)
    if not p.exists():
        return []
    # Apriamo due viste: una conserva le formule per riconoscere la riga
    # riepilogativa, l'altra restituisce i valori calcolati delle celle normali.
    wb_formula = load_workbook(p, data_only=False)
    wb_values = load_workbook(p, data_only=True)
    ws_formula = wb_formula["Dati"] if "Dati" in wb_formula.sheetnames else wb_formula.active
    ws_values = wb_values["Dati"] if "Dati" in wb_values.sheetnames else wb_values.active
    headers = [c.value for c in ws_formula[1]]
    fields = [str(h) for h in headers if h]
    definitions = field_map(profile)
    out = []
    max_row = max(ws_formula.max_row, ws_values.max_row)
    for row_idx in range(2, max_row + 1):
        if _is_generated_summary_row(ws_formula, row_idx, fields, definitions, profile):
            continue
        row = [ws_values.cell(row_idx, i + 1).value for i in range(len(headers))]
        rec = {str(headers[i]): row[i] for i in range(min(len(headers), len(row))) if headers[i]}
        if any(v not in (None, "") for v in rec.values()):
            out.append(rec)
    return out


def merge_records(existing: list[dict[str, Any]], new: list[dict[str, Any]], profile: dict[str, Any]) -> list[dict[str, Any]]:
    out = [dict(x) for x in existing]
    index: dict[tuple[Any, ...], dict[str, Any]] = {}
    for rec in out:
        k = record_key(rec, profile)
        if k is not None:
            index[k] = rec
    for rec in new:
        k = record_key(rec, profile)
        if k is not None and k in index:
            target = index[k]
            for field, value in rec.items():
                if value not in (None, ""):
                    target[field] = value
            compute_fields(target, profile)
        else:
            copy = dict(rec)
            compute_fields(copy, profile)
            out.append(copy)
            if k is not None:
                index[k] = copy
    return out



def _write_totals_table(wb, records: list[dict[str, Any]], profile: dict[str, Any]) -> None:
    cfg = profile.get("totals_table", {}) or {}
    rows = [str(x) for x in cfg.get("row_fields", []) if x]
    cols = [str(x) for x in cfg.get("column_fields", []) if x]
    values = [str(x) for x in cfg.get("value_fields", []) if x]
    date_field = str(cfg.get("date_field", "") or "")
    group_year = bool(cfg.get("group_year")); group_month = bool(cfg.get("group_month"))
    if not (rows or cols or values or (date_field and (group_year or group_month))):
        return
    ws = wb.create_sheet("Tabella totali")
    row_dims = list(rows)
    if date_field and group_year: row_dims.append(f"{date_field} - Anno")
    if date_field and group_month: row_dims.append(f"{date_field} - Mese")
    def date_parts(rec):
        v=rec.get(date_field) if date_field else None
        if isinstance(v, datetime): d=v.date()
        elif isinstance(v,date): d=v
        else:
            try: d=date.fromisoformat(str(v)[:10])
            except Exception: d=None
        return (d.year if d and group_year else "", d.month if d and group_month else "")
    def rkey(rec):
        out=[rec.get(x,"") for x in rows]; y,m=date_parts(rec)
        if group_year: out.append(y)
        if group_month: out.append(m)
        return tuple(out)
    def ckey(rec): return tuple(rec.get(x,"") for x in cols)
    rowkeys=sorted({rkey(r) for r in records}, key=lambda x: tuple(str(v) for v in x)) or [tuple()]
    colkeys=sorted({ckey(r) for r in records}, key=lambda x: tuple(str(v) for v in x)) or [tuple()]
    metrics=values or ["Conteggio"]
    headers=row_dims[:]
    for ck in colkeys:
        prefix=" / ".join(str(v) for v in ck if v not in (None,""))
        for metric in metrics: headers.append((prefix+" / " if prefix else "")+metric)
    ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True); cell.fill=PatternFill("solid",fgColor="D9EAF7")
    for rk in rowkeys:
        line=list(rk)
        for ck in colkeys:
            subset=[r for r in records if rkey(r)==rk and ckey(r)==ck]
            for metric in metrics:
                if metric=="Conteggio": val=len(subset)
                else:
                    nums=[_strict_number(r.get(metric)) for r in subset]; val=sum(x for x in nums if x is not None)
                line.append(val)
        ws.append(line)
    ws.freeze_panes="A2"
    for i in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(i)].width=min(40,max(12,max(len(str(ws.cell(r,i).value or "")) for r in range(1,ws.max_row+1))+2))


def build_period_summary_formula(
    item: dict[str, Any],
    fields: list[str],
    data_last_row: int,
) -> str:
    """Genera SOMMA/CONTA condizionata da una finestra temporale scelta dall'utente."""
    operation = str(item.get("operation", "sum") or "sum").lower()
    value_field = str(item.get("field", "") or "")
    date_field = str(item.get("date_field", "") or "")
    period = str(item.get("period", "month") or "month").lower()
    period_mode = str(item.get("period_mode", "rolling") or "rolling").lower()
    if value_field not in fields or date_field not in fields:
        return "=0"
    if data_last_row < 2:
        return "=0"
    value_col = get_column_letter(fields.index(value_field) + 1)
    date_col = get_column_letter(fields.index(date_field) + 1)
    value_range = f"{value_col}2:{value_col}{data_last_row}"
    date_range = f"{date_col}2:{date_col}{data_last_row}"

    if period_mode == "calendar":
        if period == "year":
            start_expr = "DATE(YEAR(TODAY()),1,1)"
            end_expr = "DATE(YEAR(TODAY()),12,31)"
        else:
            start_expr = "EOMONTH(TODAY(),-1)+1"
            end_expr = "EOMONTH(TODAY(),0)"
    else:
        days = 365 if period == "year" else 30
        start_expr = f"TODAY()-{days}"
        end_expr = "TODAY()"

    if operation == "count":
        return (
            f'=COUNTIFS({date_range},">="&{start_expr},{date_range},"<="&{end_expr},'
            f'{value_range},"<>")'
        )
    return (
        f'=SUMIFS({value_range},{date_range},">="&{start_expr},'
        f'{date_range},"<="&{end_expr})'
    )


def summary_formula_display(item: dict[str, Any]) -> str:
    if str(item.get("kind", "") or "") != "period_aggregate":
        return str(item.get("formula", "") or "")
    op = "Somma" if str(item.get("operation", "sum") or "sum") == "sum" else "Conta"
    period = "anno" if str(item.get("period", "month") or "month") == "year" else "mese"
    mode = "finestra mobile" if str(item.get("period_mode", "rolling") or "rolling") == "rolling" else "periodo di calendario"
    return f"{op} {item.get('field','')} se {item.get('date_field','')} è nell'ultimo {period} ({mode})"

def field_decimals(spec: dict[str, Any]) -> int:
    try:
        return max(0, min(10, int(spec.get("decimals", 2))))
    except (TypeError, ValueError):
        return 2


def excel_number_format(spec: dict[str, Any]) -> str:
    decimals = field_decimals(spec)
    decimal_part = "" if decimals == 0 else "." + ("0" * decimals)
    if str(spec.get("format", "") or "") == "currency_eur":
        return "#,##0" + decimal_part + " [$€-it-IT]"
    return "#,##0" + decimal_part


def write_excel(path: str | Path, records: list[dict[str, Any]], profile: dict[str, Any]) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Dati"
    fields = field_names(profile)
    definitions = field_map(profile)
    ws.append(fields)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for rec in records:
        ws.append([excel_value(rec.get(f), definitions.get(f, {"name": f, "type": "text"})) for f in fields])

    data_last_row = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=data_last_row):
        for idx, cell in enumerate(row):
            field = fields[idx]
            spec = definitions.get(field, {"type": "text"})
            kind = normalize_field_type(spec.get("type"))
            if kind == "number" and isinstance(cell.value, (int, float)):
                cell.number_format = excel_number_format(spec)
            elif kind == "date" and isinstance(cell.value, (date, datetime)):
                cell.number_format = "dd/mm/yyyy"
            elif kind == "datetime" and isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm:ss"

    # Una sola riga riepilogativa, immediatamente dopo l'ultimo record. Ogni
    # campo può decidere indipendentemente se scrivere SOMMA o CONTA.
    formula_fields = [
        (
            i,
            f,
            str(definitions.get(f, {}).get("final_formula", "") or "").lower(),
            str(definitions.get(f, {}).get("summary_source_field", "") or f),
        )
        for i, f in enumerate(fields, 1)
        if str(definitions.get(f, {}).get("final_formula", "") or "").lower() in {"sum", "count"}
    ]
    summary_row: int | None = None
    alias_cells: dict[str, str] = {}
    if formula_fields:
        summary_row = data_last_row + 1
        for col_idx, field_name, formula_kind, source_field in formula_fields:
            col = get_column_letter(col_idx)
            try:
                source_col_idx = fields.index(source_field) + 1
            except ValueError:
                source_col_idx = col_idx
            source_col = get_column_letter(source_col_idx)
            if data_last_row >= 2:
                if formula_kind == "sum":
                    formula = f"=SUM({source_col}2:{source_col}{data_last_row})"
                else:
                    formula = f"=COUNTA({source_col}2:{source_col}{data_last_row})"
            else:
                formula = "=0"
            cell = ws.cell(summary_row, col_idx, formula)
            cell.font = Font(bold=True)
            spec = definitions.get(field_name, {"type": "text"})
            if formula_kind == "sum" and normalize_field_type(spec.get("type")) == "number":
                cell.number_format = excel_number_format(spec)
            else:
                cell.number_format = "0"
            alias_name = ("somma_" if formula_kind == "sum" else "conta_") + field_name
            alias_cells[alias_name] = f"{col}{summary_row}"

    # Formule riepilogative nominate. Sono scritte in righe successive con il
    # nome in colonna A e il risultato/formula in colonna B. Nella definizione
    # l'utente usa alias leggibili, ad esempio:
    # =somma_importo/somma_durata_giorni
    custom_formulas = [
        item for item in profile.get("summary_formulas", []) or []
        if isinstance(item, dict) and str(item.get("name", "") or "").strip()
        and (
            str(item.get("formula", "") or "").strip()
            or str(item.get("kind", "") or "") == "period_aggregate"
        )
    ]
    if custom_formulas:
        first_custom_row = (summary_row + 1) if summary_row is not None else (data_last_row + 1)
        summary_alias_names = summary_formula_aliases(profile)
        summary_rows = {
            str(item.get("name", "") or "").strip(): first_custom_row + offset
            for offset, item in enumerate(custom_formulas)
        }
        for alias, formula_name in summary_alias_names.items():
            if formula_name in summary_rows:
                alias_cells[alias] = f"B{summary_rows[formula_name]}"

        for offset, item in enumerate(custom_formulas):
            row_idx = first_custom_row + offset
            name = str(item.get("name", "") or "").strip()
            if str(item.get("kind", "") or "") == "period_aggregate":
                expanded = build_period_summary_formula(item, fields, data_last_row)
            else:
                expanded = expand_summary_formula(str(item.get("formula", "") or ""), profile, alias_cells)
            name_cell = ws.cell(row_idx, 1, name)
            value_cell = ws.cell(row_idx, 2, expanded)
            name_cell.font = Font(bold=True)
            value_cell.font = Font(bold=True)
            if str(item.get("kind", "") or "") == "period_aggregate":
                source_spec = definitions.get(str(item.get("field", "") or ""), {})
                if normalize_field_type(source_spec.get("type")) == "number":
                    value_cell.number_format = excel_number_format(source_spec)

    ws.freeze_panes = "A2"
    # Il filtro comprende solo intestazione + dati, non la riga di formule.
    if fields:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{max(1, data_last_row)}"
    for i, f in enumerate(fields, 1):
        values = [str(f)] + [str(ws.cell(r, i).value or "") for r in range(2, data_last_row + 1)]
        width = min(45, max(12, max(len(v) for v in values) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width

    _write_totals_table(wb, records, profile)

    # Chiede a Excel/LibreOffice di ricalcolare le formule all'apertura.
    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass
    wb.save(p)



def read_excel_home_summary_values(
    path: str | Path,
    profile: dict[str, Any],
    selected_fields: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Legge i valori riepilogativi della home senza avviare Microsoft Excel.

    openpyxl legge prima gli eventuali valori già memorizzati nel file. Se la
    cache delle formule non è disponibile (caso normale subito dopo un
    salvataggio openpyxl), SOMMA/CONTA e le formule riepilogative vengono
    valutate direttamente sui dati del foglio. Questo evita di creare una
    seconda istanza COM di Excel solo per aggiornare la home, eliminando le
    eccezioni RPC 0x80010108/0x800706be osservate durante test e build.
    """
    import ast
    from datetime import date as _date, datetime as _datetime, timedelta as _timedelta

    p = Path(path)
    if not p.exists():
        return []
    selected = [str(x) for x in (selected_fields if selected_fields is not None else profile.get("home_summary_fields", [])) if x][:6]
    if not selected:
        return []
    definitions = field_map(profile)
    computed_names = {str(x.get("field", "") or "") for x in profile.get("computed_fields", []) or [] if isinstance(x, dict)}
    summary_defs = {str(x.get("name", "") or ""): x for x in profile.get("summary_formulas", []) or [] if isinstance(x, dict) and x.get("name")}

    # Due aperture openpyxl: una per i valori cached, una per formule/dati.
    wb_values = load_workbook(p, data_only=True, read_only=False)
    wb_formulas = load_workbook(p, data_only=False, read_only=False)
    try:
        ws = wb_values["Dati"] if "Dati" in wb_values.sheetnames else wb_values.active
        wsf = wb_formulas["Dati"] if "Dati" in wb_formulas.sheetnames else wb_formulas.active
        headers = [str(c.value or "") for c in ws[1]]
        if "_EgoMeta" in wb_values.sheetnames:
            try: data_last_row = int(wb_values["_EgoMeta"]["A1"].value or 1)
            except Exception: data_last_row = max(1, ws.max_row - 1)
        else:
            data_last_row = max(1, ws.max_row - 1)
        summary_row = data_last_row + 1

        def _column_values(name):
            if name not in headers: return []
            ci=headers.index(name)+1
            return [ws.cell(r,ci).value for r in range(2,data_last_row+1)]

        def _number(v):
            if isinstance(v,bool): return float(int(v))
            if isinstance(v,(int,float)): return float(v)
            if isinstance(v,str):
                s=v.strip().replace('€','').replace(' ','')
                if not s: return None
                if ',' in s: s=s.replace('.','').replace(',','.')
                try:return float(s)
                except Exception:return None
            return None

        aliases={}
        for name,spec in definitions.items():
            kind=str(spec.get('final_formula','') or '').lower()
            vals=_column_values(name)
            if kind=='sum': aliases['somma_'+name]=sum(x for x in (_number(v) for v in vals) if x is not None)
            elif kind=='count': aliases['conta_'+name]=sum(1 for v in vals if v not in (None,''))
        summary_alias_names = summary_formula_aliases(profile)

        def _safe_arithmetic(expr, extra_aliases=None):
            text=str(expr or '').strip()
            if text.startswith('='): text=text[1:]
            merged_aliases = dict(aliases)
            merged_aliases.update(extra_aliases or {})
            for alias,value in sorted(merged_aliases.items(),key=lambda x:-len(x[0])):
                text=re.sub(r'\b'+re.escape(alias)+r'\b',str(value),text,flags=re.I)
            node=ast.parse(text,mode='eval')
            allowed=(ast.Expression,ast.BinOp,ast.UnaryOp,ast.Constant,ast.Add,ast.Sub,ast.Mult,ast.Div,ast.Pow,ast.Mod,ast.USub,ast.UAdd)
            if any(not isinstance(n,allowed) for n in ast.walk(node)): raise ValueError('Formula non aritmetica')
            return eval(compile(node,'<summary>','eval'),{'__builtins__':{}},{})

        def _as_date(v):
            if isinstance(v,_datetime): return v.date()
            if isinstance(v,_date): return v
            if isinstance(v,str):
                for fmt in ('%Y-%m-%d','%d/%m/%Y','%d-%m-%Y'):
                    try:return _datetime.strptime(v.strip()[:10],fmt).date()
                    except Exception: pass
            return None

        def _period_value(item):
            vals=_column_values(str(item.get('field','') or ''))
            dates=_column_values(str(item.get('date_field','') or ''))
            today=_date.today(); period=str(item.get('period','month') or 'month').lower(); mode=str(item.get('period_mode','rolling') or 'rolling').lower()
            if mode=='calendar':
                if period=='year': start=_date(today.year,1,1); end=_date(today.year,12,31)
                else:
                    start=_date(today.year,today.month,1)
                    end=(_date(today.year+1,1,1) if today.month==12 else _date(today.year,today.month+1,1))-_timedelta(days=1)
            else:
                start=today-_timedelta(days=365 if period=='year' else 30); end=today
            picked=[v for v,d in zip(vals,dates) if (lambda dd: dd is not None and start<=dd<=end)(_as_date(d))]
            if str(item.get('operation','sum') or 'sum').lower()=='count': return sum(1 for v in picked if v not in (None,''))
            return sum(x for x in (_number(v) for v in picked) if x is not None)

        _summary_cache = {}
        _summary_stack = set()

        def _summary_value_by_name(name):
            if name in _summary_cache:
                return _summary_cache[name]
            if name in _summary_stack:
                raise ValueError(f"Riferimento circolare tra formule riepilogative: {name}")
            item = summary_defs.get(name)
            if not item:
                raise ValueError(f"Formula riepilogativa non definita: {name}")
            _summary_stack.add(name)
            try:
                if str(item.get('kind','') or '') == 'period_aggregate':
                    value = _period_value(item)
                else:
                    dependencies = {}
                    formula_text = str(item.get('formula','') or '')
                    for alias, dep_name in summary_alias_names.items():
                        if re.search(r'(?<!\w)'+re.escape(alias)+r'(?!\w)', formula_text, flags=re.I):
                            dependencies[alias] = _summary_value_by_name(dep_name)
                    value = _safe_arithmetic(formula_text, dependencies)
                _summary_cache[name] = value
                return value
            finally:
                _summary_stack.discard(name)

        result=[]
        for ref in selected:
            source_kind,_,raw_name=str(ref).partition(':')
            if not raw_name: source_kind,raw_name='field',str(ref)
            if source_kind=='summary':
                item=summary_defs.get(raw_name,{})
                value=None
                # Preferisci un valore cached reale, se esiste.
                for r in range(summary_row,ws.max_row+1):
                    if str(ws.cell(r,1).value or '')==raw_name:
                        value=ws.cell(r,2).value; break
                if value is None:
                    try:
                        value=_summary_value_by_name(raw_name)
                    except Exception: value=None
                result.append({'field':ref,'label':raw_name,'kind':'summary','value':value,'format':''})
                continue
            if raw_name not in headers: continue
            spec=definitions.get(raw_name,{}); ci=headers.index(raw_name)+1; kind=str(spec.get('final_formula','') or '').lower()
            if kind in {'sum','count'}:
                value=ws.cell(summary_row,ci).value
                if value is None: value=aliases.get(('somma_' if kind=='sum' else 'conta_')+raw_name)
                result.append({'field':ref,'label':raw_name,'kind':kind,'value':value,'format':str(spec.get('format','') or '')})
            elif source_kind=='computed' or raw_name in computed_names:
                value=None
                for r in range(data_last_row,1,-1):
                    candidate=ws.cell(r,ci).value
                    if candidate not in (None,''): value=candidate; break
                result.append({'field':ref,'label':raw_name,'kind':'computed','value':value,'format':str(spec.get('format','') or '')})
        return result
    finally:
        wb_values.close(); wb_formulas.close()


def _is_empty(value: Any) -> bool:
    return value in (None, "")


def _as_number(value: Any) -> float | int | None:
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    return _strict_number(value)


def apply_existing_value_policy(
    old_value: Any,
    new_value: Any,
    policy: str,
    field_spec: dict[str, Any],
) -> tuple[Any, str]:
    """Applica la politica della regola quando la cella contiene già un valore.

    Restituisce ``(valore_finale, descrizione_azione)`` per poter riportare
    esattamente l'operazione nel log di estrazione.
    """
    if _is_empty(new_value):
        return old_value, "NESSUN VALORE ESTRATTO"
    if _is_empty(old_value):
        return new_value, "SCRIVI (cella vuota)"

    policy = str(policy or "replace").strip().lower()
    if policy == "keep":
        return old_value, "NON SOSTITUIRE"
    if policy == "sum":
        if normalize_field_type(field_spec.get("type")) != "number":
            return new_value, "SOSTITUISCI (Somma non ammessa: campo non numerico)"
        old_num = _as_number(old_value)
        new_num = _as_number(new_value)
        if old_num is None or new_num is None:
            return old_value, "SOMMA NON ESEGUITA (valore non numerico)"
        total = old_num + new_num
        # Conserva un int quando possibile. Per gli importi evita gli artefatti
        # floating point (es. 650.8599999999999) arrotondando ai centesimi.
        if field_spec.get("format") == "currency_eur":
            total = round(float(total), 2)
        elif isinstance(old_num, int) and isinstance(new_num, int):
            total = int(total)
        return total, "SOMMA"
    return new_value, "SOSTITUISCI"


def _log_value(value: Any) -> str:
    if value is None:
        return "<vuoto>"
    text = str(value).replace("\r", " ").replace("\n", " ")
    return text if len(text) <= 500 else text[:497] + "..."


def _record_key_text(record: dict[str, Any], profile: dict[str, Any]) -> str:
    keys = profile.get("reconcile_keys", []) or []
    if not keys:
        return "nessuna chiave di riconciliazione"
    return ", ".join(f"{k}={_log_value(record.get(k))}" for k in keys)


def _append_mail_uid_history(existing: Any, uid: Any) -> str:
    """Aggiunge un UID alla cronologia testuale senza duplicati, preservando l'ordine."""
    try:
        uid_text = str(int(uid))
    except (TypeError, ValueError):
        uid_text = str(uid or "").strip()
    if not uid_text:
        return str(existing or "")
    raw = str(existing or "").strip()
    values = [x.strip() for x in re.split(r"[,;\s]+", raw) if x.strip()] if raw else []
    if uid_text not in values:
        values.append(uid_text)
    return ", ".join(values)


def apply_detailed_extractions(
    existing: list[dict[str, Any]],
    mail_events: list[dict[str, Any]],
    profile: dict[str, Any],
    *,
    log_path: str | Path | None = None,
    mode: str = "estrazione",
    excel_path: str | Path | None = None,
    include_header: bool = True,
    include_footer: bool = True,
) -> tuple[list[dict[str, Any]], dict[str, int], str]:
    """Applica gli eventi dettagliati in ordine e produce un audit log.

    ``mail_events`` deve essere già ordinato dalla mail più vecchia alla più
    recente. Le politiche ``replace``/``keep``/``sum`` sono memorizzate sulla
    singola regola di estrazione, perciò due regole che scrivono lo stesso
    campo possono avere comportamenti differenti.
    """
    out = [dict(x) for x in existing]
    definitions = field_map(profile)
    index: dict[tuple[Any, ...], int] = {}
    for i, rec in enumerate(out):
        key = record_key(rec, profile)
        if key is not None:
            index[key] = i

    now = datetime.now().astimezone()
    lines: list[str] = []
    if include_header:
        lines.extend([
            "=" * 100,
            f"INIZIO {mode.upper()} - {now.isoformat(timespec='seconds')}",
            f"Profilo: {profile.get('name', '')}",
            f"File Excel: {Path(excel_path).resolve() if excel_path else ''}",
            f"Mail in scansione: {len(mail_events)}",
            "Ordine elaborazione: dalla mail più vecchia alla più recente",
            "=" * 100,
        ])
    stats = {"mails": len(mail_events), "matched_mails": 0, "type_matches": 0, "field_updates": 0}

    for event in mail_events:
        mail = event.get("mail")
        types = event.get("types", []) or []
        lines.extend([
            "",
            "-" * 100,
            "MAIL ESAMINATA",
            f"UID: {getattr(mail, 'uid', '')}",
            f"Message-ID: {getattr(mail, 'message_id', '') or '<assente>'}",
            f"Data/ora mail: {getattr(mail, 'date', '')}",
            f"Mittente: {getattr(mail, 'sender', '')}",
            f"Oggetto: {getattr(mail, 'subject', '')}",
            f"Selezionata per estrazione: {'SI' if types else 'NO'}",
        ])
        if not types:
            lines.append("Nessun tipo mail ha soddisfatto tutte le proprie regole di selezione.")
            continue
        stats["matched_mails"] += 1

        for type_event in types:
            stats["type_matches"] += 1
            mail_type = type_event.get("mail_type", "")
            selection = type_event.get("selection_rules", []) or []
            updates = type_event.get("updates", []) or []
            candidate = dict(type_event.get("candidate", {}) or {})
            extracted_nonempty = any(
                u.get("field") and u.get("value") not in (None, "") for u in updates
            )
            counter_fields = [
                (field_name, spec) for field_name, spec in definitions.items()
                if str(spec.get("mail_type_counter", "") or "") == str(mail_type)
            ]

            lines.extend([
                "",
                f"TIPO MAIL MATCH: {mail_type}",
                "Regole di selezione che hanno determinato il match:",
            ])
            for sr in selection:
                if str(sr.get("source", "") or "") == "date":
                    lines.append(
                        "  [MATCH] "
                        f"{sr.get('name', '')} | origine=data | modalità=intervallo | "
                        f"da={_log_value(sr.get('valid_from', ''))} | a={_log_value(sr.get('valid_to', ''))}"
                    )
                else:
                    lines.append(
                        "  [MATCH] "
                        f"{sr.get('name', '')} | origine={sr.get('source', '')} | "
                        f"modalità={sr.get('mode', '')} | trigger={_log_value(sr.get('trigger', ''))}"
                    )
                    if sr.get("mode") == "regex" and sr.get("expanded_trigger") != sr.get("trigger"):
                        lines.append(
                            f"          regex espansa={_log_value(sr.get('expanded_trigger', ''))}"
                        )

            if not extracted_nonempty and not counter_fields:
                lines.append("Riga Excel: nessuna (le regole di estrazione non hanno prodotto valori).")
                lines.append("Regole di estrazione:")
                for upd in updates:
                    status = "MATCH" if upd.get("matched") else "NO MATCH"
                    lines.append(
                        f"  [{status}] {upd.get('rule_name', '')} -> campo={upd.get('field', '')} | "
                        f"valore={_log_value(upd.get('value'))} | errore={upd.get('error', '') or '<nessuno>'}"
                    )
                continue
            if not extracted_nonempty and counter_fields and profile.get("reconcile_keys"):
                lines.append(
                    "Riga Excel: nessuna. Il Tipo mail ha un contatore automatico, ma senza "
                    "un valore estratto non è possibile individuare la riga tramite la chiave di riconciliazione."
                )
                continue

            key = record_key(candidate, profile)
            if key is not None and key in index:
                row_index = index[key]
                target = out[row_index]
                row_status = "riga esistente riconciliata"
            else:
                target = {}
                out.append(target)
                row_index = len(out) - 1
                row_status = "nuova riga"
                if key is not None:
                    index[key] = row_index
            excel_row = row_index + 2
            lines.append(
                f"Riga Excel: {excel_row} ({row_status}) | {_record_key_text(candidate, profile)}"
            )
            # Campi tecnici UID: registrano tutte le mail che hanno creato o
            # aggiornato questa riga. L'UID viene aggiunto una sola volta.
            mail_uid = getattr(mail, "uid", None)
            for uid_field, uid_spec in definitions.items():
                if not bool(uid_spec.get("auto_mail_uids")):
                    continue
                old_uid_history = target.get(uid_field)
                new_uid_history = _append_mail_uid_history(old_uid_history, mail_uid)
                if new_uid_history != old_uid_history:
                    target[uid_field] = new_uid_history
                    stats["field_updates"] += 1
                lines.extend([
                    f"  [UID MAIL] campo={uid_field} | riga Excel={excel_row}",
                    f"           UID corrente={_log_value(mail_uid)}",
                    f"           valore prima={_log_value(old_uid_history)}",
                    f"           valore dopo={_log_value(target.get(uid_field))}",
                ])
            lines.append("Regole di estrazione e campi:")

            for upd in updates:
                rule_name = str(upd.get("rule_name", "") or "")
                field = str(upd.get("field", "") or "")
                value = upd.get("value")
                matched = bool(upd.get("matched"))
                policy = str(upd.get("policy", "replace") or "replace")
                if not field:
                    lines.append(f"  [IGNORA] {rule_name}: campo destinazione non impostato")
                    continue
                if value in (None, ""):
                    status = "MATCH SENZA VALORE" if matched else "NO MATCH"
                    lines.append(
                        f"  [{status}] regola={rule_name} | campo={field} | "
                        f"strategia={upd.get('strategy', '')} | valore=<vuoto> | "
                        f"errore={upd.get('error', '') or '<nessuno>'}"
                    )
                    continue

                spec = definitions.get(field, {"name": field, "type": "text"})
                old = target.get(field)
                final, action = apply_existing_value_policy(old, value, policy, spec)
                target[field] = final
                stats["field_updates"] += 1
                lines.extend([
                    f"  [ESTRAI] regola={rule_name} | campo={field} | riga Excel={excel_row}",
                    f"           origine={upd.get('source', '')} | strategia={upd.get('strategy', '')}",
                    f"           valore estratto={_log_value(value)}",
                    f"           politica su valore esistente={policy}",
                    f"           valore prima={_log_value(old)}",
                    f"           azione={action}",
                    f"           valore dopo={_log_value(final)}",
                ])
                if upd.get("strategy") == "regex" and upd.get("expanded_pattern") != upd.get("pattern"):
                    lines.append(
                        f"           regex espansa={_log_value(upd.get('expanded_pattern', ''))}"
                    )
                if upd.get("expanded_value_pattern") and upd.get("expanded_value_pattern") != upd.get("value_pattern"):
                    lines.append(
                        f"           regex valore espansa={_log_value(upd.get('expanded_value_pattern', ''))}"
                    )

            # Contatori automatici dei match Tipo mail: una mail incrementa di
            # una unità ogni campo associato al tipo, indipendentemente dal
            # numero di regole di estrazione che hanno prodotto valori.
            for counter_field, counter_spec in counter_fields:
                old_counter = target.get(counter_field)
                old_num = _as_number(old_counter)
                if old_num is None:
                    old_num = 0
                final_counter = int(old_num) + 1
                target[counter_field] = final_counter
                stats["field_updates"] += 1
                lines.extend([
                    f"  [CONTA TIPO MAIL] tipo={mail_type} | campo={counter_field} | riga Excel={excel_row}",
                    f"           valore prima={_log_value(old_counter)}",
                    f"           azione=INCREMENTA MATCH TIPO MAIL",
                    f"           valore dopo={final_counter}",
                ])

            explicit_fields = {
                str(upd.get("field", "") or "")
                for upd in updates
                if upd.get("field") and upd.get("value") not in (None, "")
            }
            computed_before = {
                str(comp.get("field", "") or ""): target.get(str(comp.get("field", "") or ""))
                for comp in profile.get("computed_fields", []) or []
                if isinstance(comp, dict) and comp.get("field")
            }
            compute_fields(target, profile, protected_fields=explicit_fields)
            for comp in profile.get("computed_fields", []) or []:
                if not isinstance(comp, dict):
                    continue
                computed_field = str(comp.get("field", "") or "")
                if not computed_field:
                    continue
                priority = str(comp.get("priority", "rule_wins") or "rule_wins")
                old_computed = computed_before.get(computed_field)
                new_computed = target.get(computed_field)
                if computed_field in explicit_fields and priority != "always":
                    lines.extend([
                        f"  [CAMPO CALCOLATO] campo={computed_field} | tipo={comp.get('kind', '')}",
                        "           azione=NON RICALCOLATO (priorità al valore della regola corrente)",
                        f"           valore mantenuto={_log_value(new_computed)}",
                    ])
                elif new_computed != old_computed:
                    lines.extend([
                        f"  [CAMPO CALCOLATO] campo={computed_field} | tipo={comp.get('kind', '')}",
                        f"           valore prima={_log_value(old_computed)}",
                        "           azione=RICALCOLA",
                        f"           valore dopo={_log_value(new_computed)}",
                    ])
            # La chiave può essere stata popolata durante gli aggiornamenti.
            final_key = record_key(target, profile)
            if final_key is not None:
                index[final_key] = row_index

    end = datetime.now().astimezone()
    if include_footer:
        lines.extend([
            "",
            "=" * 100,
            f"FINE {mode.upper()} - {end.isoformat(timespec='seconds')}",
            f"Mail esaminate: {stats['mails']}",
            f"Mail con almeno un tipo matchato: {stats['matched_mails']}",
            f"Match di tipo mail: {stats['type_matches']}",
            f"Aggiornamenti campo eseguiti: {stats['field_updates']}",
            f"Righe dati finali: {len(out)}",
            "=" * 100,
            "",
        ])
    log_text = "\n".join(lines)
    if log_path:
        append_extraction_log(log_path, log_text)
    return out, stats, log_text


def append_extraction_log(log_path: str | Path, text: str) -> None:
    lp = Path(log_path)
    lp.parent.mkdir(parents=True, exist_ok=True)
    with lp.open("a", encoding="utf-8") as fh:
        fh.write(text)
        if not text.endswith("\n"):
            fh.write("\n")
