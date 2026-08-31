import json, zipfile
from pathlib import Path
from openpyxl import load_workbook
from egomail.profile import ProfileManager
from egomail.profile_transfer import export_profile_package, import_profile_package
from egomail.exporter import write_excel

def test_profile_package_exports_and_imports_icon(tmp_path):
    profiles=tmp_path/'profiles'; profiles.mkdir(); (profiles/'logo.png').write_bytes(b'fakepng')
    p=profiles/'p.json'; p.write_text(json.dumps({'name':'Profilo','icon':'logo.png'}),encoding='utf-8')
    z=export_profile_package(p,tmp_path/'out',[])
    with zipfile.ZipFile(z) as f: assert 'profile_icon.png' in f.namelist()
    dest,_=import_profile_package(z,ProfileManager(tmp_path/'imported'))
    data=json.loads(dest.read_text(encoding='utf-8')); assert data['icon']; assert (dest.parent/data['icon']).exists()

def test_totals_table_sheet(tmp_path):
    profile={'fields':[{'name':'tipo','type':'text'},{'name':'data','type':'date'},{'name':'importo','type':'number'}], 'totals_table':{'row_fields':['tipo'],'column_fields':[],'value_fields':['importo'],'date_field':'data','group_year':True,'group_month':True}}
    x=tmp_path/'x.xlsx'; write_excel(x,[{'tipo':'A','data':'2026-08-01','importo':10},{'tipo':'A','data':'2026-08-02','importo':5}],profile)
    ws=load_workbook(x,data_only=False)['Tabella totali']; assert ws.cell(2,4).value==15

def test_ui_has_protocol_workspace_and_no_author_name():
    text=(Path(__file__).parents[1]/'egomail'/'ui.py').read_text(encoding='utf-8')
    assert 'Cartella workspace' in text and 'values=("IMAP",)' in text and 'Tabella totali…' in text
    assert 'Ennio Balocchi' not in text

def test_community_exposes_pro_gates():
    text=(Path(__file__).parents[1]/'egomail'/'ui.py').read_text(encoding='utf-8')
    assert 'Per questa funzionalità scaricare la versione Pro' in text
    assert 'menu.add_cascade(label="Wizard", menu=wiz)' in text
