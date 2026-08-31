import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from egomail.exporter import apply_detailed_extractions, load_existing_records, write_excel
from egomail.extraction import extract_mail_detailed, extract_rule
from egomail.mail import MailMessage
from egomail.profile import normalize_profile, rename_mail_type


class MailTypeCounterTests(unittest.TestCase):
    def test_counter_increments_once_per_matching_mail_on_same_row(self):
        profile = normalize_profile({
            "name": "Counter",
            "fields": [
                {"name": "id", "type": "text"},
                {"name": "n_promemoria", "type": "number", "mail_type_counter": "promemoria"},
            ],
            "reconcile_keys": ["id"],
            "mail_types": [{
                "name": "promemoria",
                "match_rules": [{"name": "m", "source": "subject", "mode": "contains", "trigger": "Promemoria"}],
                "rules": [
                    {"name": "id", "field": "id", "source": "body", "strategy": "regex", "pattern": r"ID=(\w+)", "group": 1},
                ],
            }],
        })
        mails = [
            MailMessage(uid=1, sender="x", subject="Promemoria 1", date="Mon, 1 Jun 2026 10:00:00 +0200", message_id="<1>", body="ID=ABC"),
            MailMessage(uid=2, sender="x", subject="Promemoria 2", date="Tue, 2 Jun 2026 10:00:00 +0200", message_id="<2>", body="ID=ABC"),
        ]
        events = [extract_mail_detailed(m, profile) for m in mails]
        records, stats, log = apply_detailed_extractions([], events, profile)
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["id"], "ABC")
        self.assertEqual(records[0]["n_promemoria"], 2)
        self.assertIn("CONTA TIPO MAIL", log)
        self.assertEqual(stats["field_updates"], 4)  # 2 id + 2 incrementi

    def test_mail_type_rename_updates_counter_reference(self):
        profile = normalize_profile({
            "fields": [{"name": "n", "type": "number", "mail_type_counter": "old"}],
            "mail_types": [{"name": "old", "match_rules": [], "rules": []}],
        })
        rename_mail_type(profile, "old", "nuovo")
        self.assertEqual(profile["fields"][0]["mail_type_counter"], "nuovo")


class NamedSummaryFormulaTests(unittest.TestCase):
    def test_named_formula_uses_symbolic_final_formula_aliases(self):
        profile = normalize_profile({
            "fields": [
                {"name": "importo", "type": "number", "final_formula": "sum"},
                {"name": "durata_giorni", "type": "number", "final_formula": "sum"},
            ],
            "summary_formulas": [
                {"name": "Media per giorno", "formula": "=somma_importo/somma_durata_giorni"},
            ],
        })
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "named.xlsx"
            write_excel(path, [
                {"importo": 100, "durata_giorni": 2},
                {"importo": 200, "durata_giorni": 3},
            ], profile)
            wb = load_workbook(path, data_only=False)
            ws = wb["Dati"]
            self.assertEqual(ws["A4"].value, "=SUM(A2:A3)")
            self.assertEqual(ws["B4"].value, "=SUM(B2:B3)")
            self.assertEqual(ws["A5"].value, "Media per giorno")
            self.assertEqual(ws["B5"].value, "=A4/B4")
            rows = load_existing_records(path, profile)
            self.assertEqual(len(rows), 2)

    def test_unknown_alias_is_rejected(self):
        profile = normalize_profile({
            "fields": [{"name": "importo", "type": "number", "final_formula": "sum"}],
            "summary_formulas": [{"name": "x", "formula": "=somma_inesistente/2"}],
        })
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(ValueError):
                write_excel(Path(td) / "bad.xlsx", [{"importo": 10}], profile)


class MailMetadataExtractionSourceTests(unittest.TestCase):
    def test_extract_from_mail_date_and_message_id(self):
        mail = MailMessage(
            uid=7,
            sender="x@example.com",
            subject="Oggetto",
            date="Sun, 25 Oct 2026 11:12:13 +0200",
            raw_date="Sun, 25 Oct 2026 11:12:13 +0200",
            message_id="<abc-123@example.com>",
            body="corpo",
        )
        date_value = extract_rule(mail, {
            "source": "date", "strategy": "regex", "pattern": r"(25 Oct 2026)", "group": 1,
        })
        id_value = extract_rule(mail, {
            "source": "message_id", "strategy": "regex", "pattern": r"<([^>]+)>", "group": 1,
        })
        self.assertEqual(date_value, "25 Oct 2026")
        date_iso = extract_rule(mail, {
            "source": "date", "strategy": "regex", "pattern": r"(.+)", "group": 1,
            "transform": "date_it_email_year",
        })
        self.assertEqual(date_iso, "2026-10-25")
        self.assertEqual(id_value, "abc-123@example.com")


if __name__ == "__main__":
    unittest.main()
