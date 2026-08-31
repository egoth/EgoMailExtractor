import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from egomail.exporter import write_excel, excel_value


class ExcelTypeTests(unittest.TestCase):
    def test_excel_cells_use_profile_types(self):
        profile = {
            'fields': [
                {'name': 'codice', 'type': 'text'},
                {'name': 'data_check_in', 'type': 'date'},
                {'name': 'ospiti', 'type': 'number'},
                {'name': 'importo', 'type': 'number', 'format': 'currency_eur'},
            ]
        }
        records = [{
            'codice': '00123',
            'data_check_in': '2026-08-16',
            'ospiti': '5',
            'importo': '395,54',
        }]
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'out.xlsx'
            write_excel(path, records, profile)
            wb = load_workbook(path, data_only=True)
            ws = wb['Dati']
            self.assertEqual(ws['A2'].value, '00123')
            self.assertEqual(ws['B2'].value.date(), date(2026, 8, 16))
            self.assertEqual(ws['C2'].value, 5.0)
            self.assertAlmostEqual(ws['D2'].value, 395.54)
            self.assertIn('€', ws['D2'].number_format)

    def test_alphanumeric_value_is_not_destroyed_if_field_is_mistakenly_number(self):
        # Regressione v1.10: HM895ZFM4S diventava 8954.
        self.assertEqual(excel_value("HM895ZFM4S", {"name": "id", "type": "number"}), "HM895ZFM4S")


if __name__ == '__main__':
    unittest.main()

class ExcelFieldOrderTests(unittest.TestCase):
    def test_profile_field_order_is_excel_column_order(self):
        profile = {
            'fields': [
                {'name': 'terzo', 'type': 'text'},
                {'name': 'primo', 'type': 'text'},
                {'name': 'secondo', 'type': 'text'},
            ]
        }
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'order.xlsx'
            write_excel(path, [{'primo': '1', 'secondo': '2', 'terzo': '3'}], profile)
            wb = load_workbook(path, data_only=True)
            ws = wb['Dati']
            self.assertEqual([ws.cell(1, c).value for c in range(1, 4)], ['terzo', 'primo', 'secondo'])


class ExcelFinalFormulaTests(unittest.TestCase):
    def test_sum_and_count_formulas_are_written_after_data(self):
        profile = {
            'fields': [
                {'name': 'nome', 'type': 'text', 'final_formula': 'count'},
                {'name': 'importo', 'type': 'number', 'format': 'currency_eur', 'final_formula': 'sum'},
                {'name': 'note', 'type': 'text'},
            ]
        }
        records = [
            {'nome': 'Anna', 'importo': 10.5, 'note': 'x'},
            {'nome': 'Luca', 'importo': 20, 'note': 'y'},
        ]
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'formule.xlsx'
            write_excel(path, records, profile)
            wb = load_workbook(path, data_only=False)
            ws = wb['Dati']
            self.assertEqual(ws['A4'].value, '=COUNTA(A2:A3)')
            self.assertEqual(ws['B4'].value, '=SUM(B2:B3)')
            self.assertIsNone(ws['C4'].value)
            self.assertEqual(ws.auto_filter.ref, 'A1:C3')
            self.assertTrue(ws['A4'].font.bold)
            self.assertTrue(ws['B4'].font.bold)
            self.assertIn('€', ws['B4'].number_format)

    def test_summary_row_is_not_loaded_as_existing_record(self):
        from egomail.exporter import load_existing_records
        profile = {
            'fields': [
                {'name': 'id', 'type': 'text'},
                {'name': 'nome', 'type': 'text', 'final_formula': 'count'},
                {'name': 'importo', 'type': 'number', 'final_formula': 'sum'},
            ],
            'reconcile_keys': ['id'],
        }
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'formule.xlsx'
            write_excel(path, [
                {'id': 'A', 'nome': 'Anna', 'importo': 10},
                {'id': 'B', 'nome': 'Luca', 'importo': 20},
            ], profile)
            rows = load_existing_records(path, profile)
            self.assertEqual(len(rows), 2)
            self.assertEqual([r['id'] for r in rows], ['A', 'B'])

    def test_empty_dataset_uses_zero_without_self_reference(self):
        profile = {'fields': [{'name': 'importo', 'type': 'number', 'final_formula': 'sum'}]}
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / 'vuoto.xlsx'
            write_excel(path, [], profile)
            wb = load_workbook(path, data_only=False)
            ws = wb['Dati']
            self.assertEqual(ws['A2'].value, '=0')
            self.assertEqual(ws.auto_filter.ref, 'A1:A1')
